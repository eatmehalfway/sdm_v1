"""Load the orthopaedic demo workbook into the department episode shape."""
from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKBOOK_PATH = Path(__file__).resolve().parent / "data" / "ccq_orthopaedic_demo_dataset.xlsx"

DEPARTMENT_NAME = "Orthopedic Surgery"

CLINICIANS = [
    {"id": "clin_a", "name": "Clinician A", "color": "#1769e0"},
    {"id": "clin_b", "name": "Clinician B", "color": "#08a3b7"},
    {"id": "clin_c", "name": "Clinician C", "color": "#5968cf"},
    {"id": "clin_d", "name": "Clinician D", "color": "#338bc4"},
    {"id": "clin_e", "name": "Clinician E", "color": "#1451a4"},
]

_CLINICIAN_BY_NAME = {row["name"]: row for row in CLINICIANS}

_EPISODES: list[dict[str, Any]] | None = None


def _excel_date(value: Any) -> str:
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(value))).date().isoformat()
    text = str(value or "").strip()
    return text[:10]


def _sheet_rows(workbook, name: str) -> list[dict[str, Any]]:
    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    parsed = []
    for raw in rows:
        if not raw or raw[0] in (None, ""):
            continue
        parsed.append({header[index]: raw[index] if index < len(raw) else None for index in range(len(header))})
    return parsed


def _flag(value: Any) -> bool:
    if value in (None, ""):
        return False
    return str(value).strip() not in {"0", "0.0", "false", "False"}


def load_episodes() -> list[dict[str, Any]]:
    global _EPISODES
    if _EPISODES is not None:
        return list(_EPISODES)

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        episode_rows = _sheet_rows(workbook, "Episodes")
        risk_rows = _sheet_rows(workbook, "Risk_Long")
    finally:
        workbook.close()

    risks_by_episode: dict[str, list[dict[str, Any]]] = {}
    for row in risk_rows:
        if not _flag(row.get("applicable")):
            continue
        episode_id = str(row.get("episode_id") or "").strip()
        domain = str(row.get("risk_domain") or "").strip()
        if not episode_id or not domain:
            continue
        discussed = _flag(row.get("mentioned"))
        risks_by_episode.setdefault(episode_id, []).append({
            "risk_name": domain,
            "risk_category": domain,
            "detection_status": "Discussed" if discussed else "Not detected",
        })

    episodes = []
    for row in episode_rows:
        episode_id = str(row.get("episode_id") or "").strip()
        clinician_name = str(row.get("clinician") or "").strip()
        clinician = _CLINICIAN_BY_NAME.get(clinician_name)
        if not episode_id or not clinician:
            continue
        procedure = str(row.get("intervention") or "").strip() or None
        framing = str(row.get("comparative_framing") or "").strip() or None
        items = risks_by_episode.get(episode_id, [])
        discussed = sum(1 for item in items if item["detection_status"] == "Discussed")
        episodes.append({
            "episode_id": episode_id,
            "source": "demo",
            "run_id": None,
            "department": DEPARTMENT_NAME,
            "clinician_id": clinician["id"],
            "clinician_name": clinician["name"],
            "clinician_color": clinician["color"],
            "patient_label": episode_id,
            "decision_label": procedure or "Intervention",
            "index_decision_id": f"{episode_id}_d1",
            "has_index_decision": True,
            "procedure": procedure,
            "status": "Resolved",
            "ccq_score": None,
            "decision_framing": None,
            "tradeoff_level": framing,
            "tradeoff_pct": None,
            "engagement_pct": None,
            "core_risk_pct": round(100 * discussed / len(items)) if items else None,
            "teachback": None,
            "transcribed_encounters": None,
            "total_encounters": None,
            "completeness_pct": None,
            "core_risk_items": items,
            "patient_questions": [],
            "recommendation_explained": False,
            "questions_answered_count": 0,
            "questions_total_count": 0,
            "run_date": None,
            "episode_date": _excel_date(row.get("date")),
        })
    _EPISODES = episodes
    return list(episodes)


def episode_detail(episode_id: str) -> dict[str, Any] | None:
    meta = next((episode for episode in load_episodes() if episode["episode_id"] == episode_id), None)
    if not meta:
        return None
    procedure = meta.get("procedure")
    risks = [
        {
            "risk_name": item["risk_name"],
            "risk_category": item["risk_category"],
            "detection_status": item["detection_status"],
            "quote": "",
            "turn_indices": [],
        }
        for item in meta.get("core_risk_items") or []
    ]
    decision = {
        "decision_id": f"{episode_id}_d1",
        "parent_decision_id": None,
        "decision_label": procedure or "Intervention",
        "status": "Resolved",
        "selected_option": procedure,
        "selected_intervention": procedure,
        "linked_interventions": [procedure] if procedure else [],
        "option_names": [],
        "decision_overview": {"options_considered": []},
        "informed_consent_analysis": {
            "interventions": [{
                "intervention_name": procedure,
                "core_risks": risks,
            }],
        },
        "generative_layer": {},
    }
    return {
        "episode_id": episode_id,
        "source": "demo",
        "episode": meta,
        "turns": [],
        "regions": {
            "schema_version": 2,
            "default_procedure": procedure,
            "decisions": [decision],
        },
        "schema_version": 2,
    }
